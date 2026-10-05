"""Build the advanced dark-theme AI-to-agents session deck."""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parent
OUT = ROOT / "From_AI_to_Agents_Advanced_Dark.pptx"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank = prs.slide_layouts[6]

BG = "060B18"
PANEL = "0C1628"
PANEL2 = "101F36"
LINE = "1D304A"
TEXT = "F4F8FF"
MUTED = "91A6C0"
CYAN = "37E6D5"
BLUE = "5B8CFF"
PURPLE = "9368FF"
PINK = "F472B6"
ORANGE = "FFB454"
GREEN = "43D17A"
RED = "FF6B6B"


def rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value)


def shape(slide, kind, x, y, w, h, fill, *, line=None, lw=1, transparency=0, rotation=0):
    item = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    item.fill.solid()
    item.fill.fore_color.rgb = rgb(fill)
    item.fill.transparency = transparency
    if line:
        item.line.color.rgb = rgb(line)
        item.line.width = Pt(lw)
    else:
        item.line.fill.background()
    item.rotation = rotation
    return item


def box(slide, x, y, w, h, fill=PANEL, *, line=LINE, radius=True, transparency=0, lw=1):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    return shape(slide, kind, x, y, w, h, fill, line=line, lw=lw, transparency=transparency)


def txt(slide, x, y, w, h, value, size=18, color=TEXT, bold=False, *,
        align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, font="Avenir Next", margin=0,
        italic=False):
    item = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = item.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = Inches(margin)
    frame.vertical_anchor = valign
    p = frame.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = value
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = rgb(color)
    return item


def rich(slide, x, y, w, h, runs, size=18, *, align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP):
    item = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = item.text_frame
    frame.clear(); frame.word_wrap = True; frame.vertical_anchor = valign
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
    p = frame.paragraphs[0]; p.alignment = align
    for value, color, bold in runs:
        r = p.add_run(); r.text = value; r.font.name = "Avenir Next"; r.font.size = Pt(size)
        r.font.color.rgb = rgb(color); r.font.bold = bold
    return item


def circle(slide, x, y, d, fill, label="", size=15, color=BG, *, line=None, transparency=0):
    item = shape(slide, MSO_SHAPE.OVAL, x, y, d, d, fill, line=line, transparency=transparency)
    if label:
        frame = item.text_frame; frame.clear(); frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = frame.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
        r = p.add_run(); r.text = label; r.font.name = "Avenir Next"; r.font.size = Pt(size)
        r.font.bold = True; r.font.color.rgb = rgb(color)
    return item


def connector(slide, x1, y1, x2, y2, color=LINE, width=1.4):
    item = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    item.line.color.rgb = rgb(color); item.line.width = Pt(width)
    return item


def chevron(slide, x, y, w=.32, h=.42, color=CYAN, rotation=0):
    return shape(slide, MSO_SHAPE.CHEVRON, x, y, w, h, color, transparency=5, rotation=rotation)


def pill(slide, x, y, w, label, *, fill=PANEL2, color=CYAN, line=None):
    box(slide, x, y, w, .34, fill, line=line or fill)
    txt(slide, x, y+.01, w, .24, label.upper(), 8.5, color, True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)


def canvas(slide, *, number=None, section="FROM AI TO AGENTS"):
    slide.background.fill.solid(); slide.background.fill.fore_color.rgb = rgb(BG)
    # corner accents + dot grid (kept native for reliable rendering)
    circle(slide, 12.15, -.32, .72, PANEL2)
    circle(slide, -.28, 6.82, .62, PANEL2)
    for row in range(6):
        for col in range(12):
            circle(slide, .28 + col*1.15, .32 + row*1.15, .025, LINE)
    txt(slide, .48, 7.2, 4.2, .16, section, 7.5, MUTED, True)
    if number is not None:
        txt(slide, 12.35, 7.18, .45, .18, f"{number:02}", 8, MUTED, True, align=PP_ALIGN.RIGHT)


def header(slide, title, kicker, number):
    canvas(slide, number=number)
    pill(slide, .6, .42, 1.75, kicker)
    txt(slide, .6, .93, 12.0, .62, title, 27, TEXT, True)


def node(slide, x, y, w, h, title, subtitle="", *, accent=CYAN, fill=PANEL):
    box(slide, x, y, w, h, fill, line=LINE, lw=1)
    shape(slide, MSO_SHAPE.RECTANGLE, x, y, .06, h, accent)
    txt(slide, x+.2, y+.15, w-.35, .28, title, 12.5, TEXT, True)
    if subtitle:
        txt(slide, x+.2, y+.5, w-.35, h-.58, subtitle, 9.5, MUTED)


def metric(slide, x, y, value, label, *, color=CYAN, w=2.15):
    box(slide, x, y, w, 1.18, PANEL, line=LINE)
    txt(slide, x+.2, y+.18, w-.4, .42, value, 25, color, True)
    txt(slide, x+.2, y+.72, w-.4, .22, label.upper(), 8.5, MUTED, True)


def section_band(slide, x, y, w, title, subtitle, color):
    box(slide, x, y, w, 1.08, PANEL, line=LINE)
    circle(slide, x+.22, y+.25, .56, color, "", transparency=5)
    txt(slide, x+.98, y+.2, w-1.2, .26, title, 14, TEXT, True)
    txt(slide, x+.98, y+.57, w-1.2, .23, subtitle, 10, MUTED)


# 01 — Hero
s = prs.slides.add_slide(blank); canvas(s, number=1)
shape(s, MSO_SHAPE.RECTANGLE, 0, 0, 13.333, .1, CYAN)
pill(s, .7, .62, 2.28, "Intelligent systems", fill=PANEL2)
txt(s, .7, 1.32, 7.8, .65, "FROM AI", 38, MUTED, True)
txt(s, .7, 2.05, 8.3, .8, "TO AGENTS", 48, TEXT, True)
txt(s, .72, 3.1, 7.8, .72, "Building intelligent systems from single to multi-agent", 20, CYAN, True)
txt(s, .72, 4.2, 6.75, .8, "Reason. Use tools. Take action. Collaborate—with control.", 16, MUTED)
pill(s, .72, 5.55, 2.3, "Gemini API demos", fill=PANEL, color=ORANGE, line=LINE)
txt(s, .72, 6.18, 5.5, .3, "ST. PETER'S  •  ADVANCED SESSION DECK", 9, MUTED, True)
# Hero agent constellation
circle(s, 9.55, 2.55, 1.55, TEXT, "AI", 27, BG)
for x,y,d,c,label in [(8.25,1.0,.9,CYAN,"M"),(10.95,.9,1.05,PURPLE,"T"),(11.25,3.45,.82,ORANGE,"A"),(9.8,5.0,.92,BLUE,"R"),(7.65,4.2,1.05,PINK,"G")]:
    connector(s,10.32,3.32,x+d/2,y+d/2,LINE,1.5); circle(s,x,y,d,c,label,13,BG)
for x,y,d,c in [(7.7,.65,.22,BLUE),(12.2,2.2,.3,CYAN),(8.8,5.65,.18,ORANGE),(11.8,5.2,.24,PINK)]:
    circle(s,x,y,d,c,transparency=10)

# 02 — Session promise
s=prs.slides.add_slide(blank); header(s,"What this session enables","Outcomes",2)
outcomes=[("01","EVOLUTION","AI → GenAI → agents",CYAN),("02","ANATOMY","model • tools • memory",PURPLE),("03","SINGLE AGENT","design one reliable loop",BLUE),("04","MULTI-AGENT","coordinate specialists",PINK),("05","DECISION","choose the right pattern",ORANGE),("06","PRODUCTION","evaluate, govern, observe",GREEN)]
for i,(n,t,b,c) in enumerate(outcomes):
    x=.62+(i%3)*4.12; y=1.78+(i//3)*1.72
    box(s,x,y,3.78,1.42,PANEL,line=LINE)
    txt(s,x+.2,y+.18,.45,.22,n,9,c,True)
    txt(s,x+.78,y+.17,2.65,.26,t,12.5,TEXT,True)
    txt(s,x+.78,y+.62,2.65,.34,b,11,MUTED)
box(s,.62,5.45,12.05,.78,PANEL2,line=CYAN,lw=1.2)
rich(s,.92,5.68,11.45,.32,[("Deliverable: ",CYAN,True),("a reusable framework for turning an AI use case into a controlled agentic solution.",TEXT,False)],15,align=PP_ALIGN.CENTER)

# 03 — Cold open
s=prs.slides.add_slide(blank); header(s,"You already used AI today. But did any of it act for you?","Cold open",3)
items=[("FACE","unlock",BLUE),(">>","recommend",PURPLE),("MAP","route",CYAN),("MAIL","filter",ORANGE),("CAM","enhance",PINK),("CHAT","generate",GREEN)]
for i,(icon,label,c) in enumerate(items):
    x=.64+i*2.05
    circle(s,x+.46,1.75,.94,PANEL2,icon,12,c,line=c)
    txt(s,x,2.9,1.85,.26,label.upper(),10,TEXT,True,align=PP_ALIGN.CENTER)
connector(s,.76,3.58,12.55,3.58,LINE,2)
for x,c in [(1.1,BLUE),(3.15,PURPLE),(5.2,CYAN),(7.25,ORANGE),(9.3,PINK),(11.35,GREEN)]: circle(s,x,3.45,.26,c)
box(s,1.25,4.45,10.83,1.13,PANEL,line=LINE)
txt(s,1.62,4.74,10.08,.31,"AI recognizes patterns. Agents pursue outcomes.",22,TEXT,True,align=PP_ALIGN.CENTER)
txt(s,2.0,5.98,9.33,.3,"The shift is from producing an answer to managing a process.",13,CYAN,True,align=PP_ALIGN.CENTER)

# 04 — Maturity spectrum
s=prs.slides.add_slide(blank); header(s,"The capability ladder: prediction to coordinated action","Evolution",4)
stages=[("1","AI","predict","spam • routes",MUTED),(
"2","GENERATIVE AI","create","text • code",BLUE),("3","LLM","reason","language interface",PURPLE),("4","AI AGENT","act","tools • feedback",CYAN),("5","MULTI-AGENT","collaborate","specialist team",ORANGE)]
for i,(n,t,v,ex,c) in enumerate(stages):
    x=.54+i*2.55; y=4.55-i*.58
    box(s,x,y,2.17,1.3,PANEL,line=c,lw=1.2)
    txt(s,x+.16,y+.14,.35,.22,n,9,c,True)
    txt(s,x+.16,y+.45,1.86,.25,t,11.5,TEXT,True)
    txt(s,x+.16,y+.82,1.86,.2,f"{v}  /  {ex}",8.7,MUTED)
    if i<4: chevron(s,x+2.23,y+.43,.26,.35,c)
txt(s,.64,6.45,12.0,.28,"↑ capability     ↑ autonomy     ↑ coordination cost     ↑ need for control",12,MUTED,True,align=PP_ALIGN.CENTER)

# 05 — Agency boundary
s=prs.slides.add_slide(blank); header(s,"The agency boundary: answer versus outcome","Core distinction",5)
box(s,.65,1.62,5.62,4.75,PANEL,line=BLUE,lw=1.2)
pill(s,.92,1.9,1.62,"Chatbot",fill=PANEL2,color=BLUE)
txt(s,.94,2.52,4.95,.57,"“Which courses should I learn for this job?”",18,TEXT,True)
txt(s,.94,3.48,4.9,.7,"Explains options and returns recommendations.",14,MUTED)
section_band(s,.94,4.55,4.95,"OUTPUT","text response",BLUE)
pill(s,.94,5.84,1.9,"Generate",fill=PANEL2,color=BLUE)
box(s,7.03,1.62,5.62,4.75,PANEL,line=CYAN,lw=1.2)
pill(s,7.3,1.9,1.62,"Agent",fill=PANEL2,color=CYAN)
txt(s,7.32,2.52,4.9,.57,"“Prepare me for this role in 14 days.”",18,TEXT,True)
txt(s,7.32,3.48,4.9,.7,"Analyses evidence, finds gaps, creates tasks and schedules work.",14,MUTED)
section_band(s,7.32,4.55,4.95,"OUTCOME","plan + calendar action",CYAN)
pill(s,7.32,5.84,2.25,"Decide + act",fill=PANEL2,color=CYAN)

# 06 — Agent anatomy
s=prs.slides.add_slide(blank); header(s,"An agent is an engineered system—not a model with a new name","Agent anatomy",6)
circle(s,5.15,2.15,3.0,PANEL2,"AGENT",25,TEXT,line=CYAN)
circle(s,5.78,2.78,1.75,TEXT,"GEMINI",16,BG)
parts=[(.78,1.62,"GOAL","success criteria",CYAN),(1.0,4.62,"MEMORY","state + context",PURPLE),(9.95,1.62,"TOOLS","APIs + functions",ORANGE),(9.75,4.62,"POLICY","permissions + limits",RED)]
for x,y,t,b,c in parts:
    node(s,x,y,2.58,1.0,t,b,accent=c)
    connector(s,x+2.58 if x<5 else 8.15,y+.5,5.15 if x<5 else x,y+.5,c,1.4)
pill(s,5.1,5.76,3.12,"Observability closes the loop",fill=PANEL,color=GREEN,line=LINE)

# 07 — Agent loop
s=prs.slides.add_slide(blank); header(s,"The agent loop converts intention into controlled action","Execution loop",7)
cx,cy=6.65,3.72
circle(s,5.73,2.8,1.84,TEXT,"GOAL",18,BG)
steps=[(2.0,1.6,"1","PERCEIVE","goal + context",BLUE),(8.95,1.6,"2","PLAN","next best step",PURPLE),(9.25,4.55,"3","ACT","call a tool",ORANGE),(2.0,4.55,"4","OBSERVE","inspect result",GREEN)]
for x,y,n,t,b,c in steps:
    box(s,x,y,2.18,1.14,PANEL,line=c)
    circle(s,x+.18,y+.28,.48,c,n,11,BG)
    txt(s,x+.82,y+.21,1.16,.22,t,11,TEXT,True)
    txt(s,x+.82,y+.58,1.2,.22,b,8.7,MUTED)
connector(s,4.18,2.16,5.73,3.35,LINE,1.6); connector(s,7.57,3.35,8.95,2.16,LINE,1.6)
connector(s,10.04,2.74,10.3,4.55,LINE,1.6); connector(s,9.25,5.1,7.57,4.08,LINE,1.6)
connector(s,5.73,4.08,4.18,5.1,LINE,1.6); connector(s,2.98,4.55,3.05,2.74,LINE,1.6)
box(s,3.75,6.18,5.8,.5,PANEL2,line=LINE)
txt(s,3.95,6.31,5.4,.2,"STOP: success • approval • limit • escalation",10,CYAN,True,align=PP_ALIGN.CENTER)

# 08 — Function calling sequence
s=prs.slides.add_slide(blank); header(s,"Gemini proposes the action; your application executes it","Tool calling",8)
lanes=[("USER",.75,BLUE),("APP / ORCHESTRATOR",3.0,CYAN),("GEMINI API",6.45,PURPLE),("TOOL / SYSTEM",9.75,ORANGE)]
for label,x,c in lanes:
    pill(s,x,1.55,2.15,label,fill=PANEL2,color=c,line=LINE)
    connector(s,x+1.08,2.05,x+1.08,6.35,LINE,1)
events=[(1.85,3.15,"goal + constraints",BLUE),(4.05,6.55,"prompt + function schema",CYAN),(7.55,9.86,"function(name, args)",PURPLE),(9.86,7.55,"validated result",ORANGE),(6.55,4.05,"grounded final response",GREEN),(3.15,1.85,"answer / approval request",CYAN)]
ys=[2.45,3.15,3.85,4.55,5.25,5.95]
for (x1,x2,label,c),y in zip(events,ys):
    connector(s,x1,y,x2,y,c,2); chevron(s,x2-.16 if x2>x1 else x2-.05,y-.16,.24,.3,c,0 if x2>x1 else 180)
    txt(s,min(x1,x2)+.15,y-.36,abs(x2-x1)-.3,.22,label,9,c,True,align=PP_ALIGN.CENTER)
txt(s,.75,6.68,12.0,.18,"Security boundary: validate arguments, permissions and side effects before execution.",9,MUTED,True,align=PP_ALIGN.CENTER)

# 09 — Memory
s=prs.slides.add_slide(blank); header(s,"Memory is not one thing: separate state by purpose","Memory design",9)
layers=[("WORKING MEMORY","current turn • active plan","small + temporary",CYAN,1.65),(
"SESSION STATE","conversation • task progress","scoped + resumable",BLUE,2.65),("LONG-TERM MEMORY","preferences • prior outcomes","consented + curated",PURPLE,3.65),("KNOWLEDGE / RAG","documents • policies • evidence","retrieved + cited",ORANGE,4.65)]
for t,b,rule,c,y in layers:
    box(s,1.05,y,8.1,.72,PANEL,line=c)
    txt(s,1.32,y+.18,2.2,.23,t,10.5,c,True)
    txt(s,3.6,y+.18,2.85,.23,b,10,TEXT)
    pill(s,6.78,y+.17,1.95,rule,fill=PANEL2,color=MUTED)
box(s,9.75,1.65,2.55,3.72,PANEL2,line=LINE)
txt(s,10.05,1.95,1.95,.28,"DESIGN RULE",10,CYAN,True,align=PP_ALIGN.CENTER)
txt(s,10.08,2.63,1.9,1.55,"Store only what improves the next decision.",20,TEXT,True,align=PP_ALIGN.CENTER,valign=MSO_ANCHOR.MIDDLE)
txt(s,10.08,4.72,1.9,.35,"Not every chat belongs in memory.",9,MUTED,align=PP_ALIGN.CENTER)

# 10 — Previous deck real-world examples
s=prs.slides.add_slide(blank); header(s,"From everyday requests to agentic workflows","Real-world examples",10)
examples=[("FOOD DELIVERY","search → compare → approve → order",ORANGE),("TRAVEL","flights → hotels → budget → itinerary",CYAN),("RIDE COORDINATION","driver → route → pricing → safety",BLUE),("CUSTOMER SUPPORT","order → delivery → policy → refund",PURPLE),("SOFTWARE","design → code → test → review",GREEN),("CAREER RESEARCH","market → skills → roadmap → decision",PINK)]
for i,(t,b,c) in enumerate(examples):
    x=.62+(i%3)*4.12; y=1.62+(i//3)*1.7
    box(s,x,y,3.78,1.38,PANEL,line=LINE)
    circle(s,x+.25,y+.29,.52,c,str(i+1),11,BG)
    txt(s,x+.95,y+.22,2.5,.25,t,11,TEXT,True)
    txt(s,x+.95,y+.65,2.5,.32,b,10,MUTED)
box(s,1.25,5.35,10.83,.84,PANEL2,line=ORANGE)
rich(s,1.55,5.59,10.2,.3,[("Design note: ",ORANGE,True),("these are reference architectures—not claims that every named platform uses LLM agents internally.",TEXT,False)],11.5,align=PP_ALIGN.CENTER)

# 11 — Travel single vs multi
s=prs.slides.add_slide(blank); header(s,"The same travel goal, two very different architectures","Travel example",11)
box(s,.62,1.58,5.75,4.9,PANEL,line=CYAN)
pill(s,.92,1.9,1.85,"Single agent",fill=PANEL2,color=CYAN)
node(s,2.18,2.6,2.55,.88,"TRAVEL AGENT","one context • many tools",accent=CYAN)
for i,(t,c) in enumerate([("SEARCH",BLUE),("BUDGET",GREEN),("PLANNING",PURPLE)]):
    x=1.0+i*1.65; node(s,x,4.1,1.38,.7,t,"tool",accent=c); connector(s,3.45,3.48,x+.69,4.1,LINE,1.1)
txt(s,.92,5.35,4.95,.55,"Best when one agent can reliably manage the goal and toolset.",11,MUTED)
box(s,6.95,1.58,5.75,4.9,PANEL,line=ORANGE)
pill(s,7.25,1.9,2.0,"Multi-agent",fill=PANEL2,color=ORANGE)
node(s,8.53,2.43,2.55,.76,"TRAVEL MANAGER","delegate + integrate",accent=ORANGE)
for i,(t,c) in enumerate([("FLIGHT",BLUE),("HOTEL",PURPLE),("ACTIVITY",CYAN),("BUDGET",GREEN)]):
    x=7.2+(i%2)*2.65; y=3.65+(i//2)*1.02
    node(s,x,y,2.25,.66,t,"specialist",accent=c); connector(s,9.8,3.19,x+1.12,y,LINE,1.0)
txt(s,7.25,5.73,4.95,.4,"Best when subtasks need distinct expertise or parallel work.",11,MUTED)

# 12 — Software and support teams
s=prs.slides.add_slide(blank); header(s,"Specialization becomes intuitive when agents resemble a team","Team analogy",12)
box(s,.62,1.55,7.35,4.98,PANEL,line=LINE)
pill(s,.92,1.87,2.2,"Software delivery",fill=PANEL2,color=GREEN)
node(s,2.95,2.48,2.65,.72,"PROJECT MANAGER","requirements + coordination",accent=GREEN)
team=[(.92,3.75,"DEVELOPER",BLUE),(2.67,3.75,"UI AGENT",PINK),(4.42,3.75,"DB AGENT",PURPLE),(6.17,3.75,"TEST AGENT",ORANGE)]
for x,y,t,c in team:
    node(s,x,y,1.4,.72,t,"specialist",accent=c); connector(s,4.28,3.2,x+.7,y,LINE,1)
node(s,2.95,5.05,2.65,.72,"REVIEW AGENT","quality gate",accent=CYAN)
for x,*_ in team: connector(s,x+.7,4.47,4.28,5.05,LINE,1)
box(s,8.45,1.55,4.22,4.98,PANEL2,line=LINE)
pill(s,8.78,1.87,2.12,"Support workflow",fill=PANEL,color=PURPLE)
for i,(t,b,c) in enumerate([("ORDER","status",BLUE),("DELIVERY","tracking",CYAN),("POLICY","eligibility",PURPLE),("REFUND","approved action",ORANGE)]):
    y=2.58+i*.78; node(s,8.8,y,3.45,.58,t,b,accent=c)
    if i<3: chevron(s,10.38,y+.58,.28,.23,c,90)
txt(s,8.8,5.9,3.45,.24,"Clear roles make failures easier to isolate.",9.5,MUTED,True)

# 13 — Single agent architecture
s=prs.slides.add_slide(blank); header(s,"Single-agent pattern: one decision-maker, bounded capabilities","Architecture",13)
node(s,.65,2.72,1.65,.92,"STUDENT","resume + goal",accent=BLUE)
chevron(s,2.5,2.98,.34,.38,BLUE)
box(s,3.02,1.83,3.05,2.72,PANEL2,line=CYAN,lw=1.5)
pill(s,3.38,2.16,2.32,"PLACEMENT AGENT",fill=PANEL,color=CYAN)
txt(s,3.4,2.88,2.26,.55,"analyse → prioritize → plan → schedule",14,TEXT,True,align=PP_ALIGN.CENTER)
pill(s,3.65,3.84,1.8,"Gemini API",fill=TEXT,color=BG)
tools=[(7.1,1.5,"RESUME","evidence",BLUE),(9.55,1.5,"JOB","requirements",PURPLE),(7.1,3.25,"CALENDAR","action",ORANGE),(9.55,3.25,"STATE","progress",GREEN)]
for x,y,t,b,c in tools: node(s,x,y,1.85,.92,t,b,accent=c)
for x,y,*_ in tools: connector(s,6.07,3.18,x,y+.46,LINE,1.4)
node(s,8.25,5.15,3.15,.85,"PLAN + CALENDAR","human approves commitments",accent=CYAN)
connector(s,6.07,3.5,9.82,5.15,CYAN,1.8)
txt(s,.75,5.72,5.6,.5,"Use when ownership is clear and the task fits one coherent context.",13,MUTED)

# 14 — Demo 1 dashboard
s=prs.slides.add_slide(blank); header(s,"Demo 1 — Resume ↔ job gap analysis","LLM call",14)
box(s,.6,1.55,3.3,4.95,PANEL,line=LINE)
pill(s,.9,1.84,1.28,"Input",fill=PANEL2,color=BLUE)
txt(s,.9,2.38,2.65,.28,"ARJUN REDDY",16,TEXT,True)
txt(s,.9,2.86,2.62,.58,"Python • SQL • Flask\nML project • Git",12,MUTED)
connector(s,.9,3.72,3.58,3.72,LINE,1)
txt(s,.9,4.05,2.65,.28,"GRADUATE AI ENGINEER",13,TEXT,True)
txt(s,.9,4.5,2.62,.78,"Gemini API • agents\nDocker • CI/CD • evals",11,MUTED)
pill(s,.9,5.77,1.76,"One Gemini call",fill=PANEL2,color=PURPLE)
chevron(s,4.15,3.55,.55,.62,CYAN)
metric(s,4.92,1.62,"72%","match score",color=CYAN,w=2.1)
metric(s,7.23,1.62,"4","evidence-backed matches",color=GREEN,w=2.45)
metric(s,9.89,1.62,"4","priority gaps",color=ORANGE,w=2.1)
box(s,4.92,3.12,7.08,3.03,PANEL2,line=LINE)
txt(s,5.22,3.42,2.45,.25,"STRONG SIGNALS",10,GREEN,True)
txt(s,5.22,3.86,2.45,1.35,"• Python + SQL\n• REST API project\n• ML evaluation\n• communication",12,TEXT)
txt(s,8.04,3.42,2.45,.25,"HIRING GAPS",10,ORANGE,True)
txt(s,8.04,3.86,2.55,1.35,"• Gemini function calls\n• Docker deployment\n• automated tests\n• model observability",12,TEXT)
pill(s,10.55,5.55,1.1,"Honest",fill=PANEL,color=CYAN,line=LINE)
txt(s,4.95,6.42,6.8,.2,"demo_1_llm_call.py",9,MUTED,True,align=PP_ALIGN.RIGHT)

# 15 — Demo 2
s=prs.slides.add_slide(blank); header(s,"Demo 2 — One agent turns gaps into scheduled work","Single agent",15)
steps=[("1","ANALYSE","resume evidence",BLUE),("2","PRIORITIZE","impact × time",PURPLE),("3","PLAN","measurable tasks",CYAN),("4","ACT","create .ics",ORANGE)]
for i,(n,t,b,c) in enumerate(steps):
    x=.62+i*2.62
    box(s,x,1.7,2.22,1.25,PANEL,line=c)
    circle(s,x+.18,1.98,.48,c,n,11,BG)
    txt(s,x+.82,1.91,1.16,.22,t,10.5,TEXT,True)
    txt(s,x+.82,2.3,1.16,.22,b,8.5,MUTED)
    if i<3: chevron(s,x+2.3,2.08,.24,.3,c)
# calendar UI
box(s,1.0,3.45,7.65,2.65,PANEL2,line=LINE)
txt(s,1.3,3.72,2.6,.26,"14-DAY PREPARATION SPRINT",12,CYAN,True)
for i in range(14):
    x=1.32+(i%7)*.93; y=4.28+(i//7)*.75
    c=[PURPLE,BLUE,CYAN,ORANGE][i%4]
    box(s,x,y,.72,.54,PANEL,line=c)
    txt(s,x,y+.08,.72,.18,str(i+1),8,c,True,align=PP_ALIGN.CENTER)
    shape(s,MSO_SHAPE.RECTANGLE,x+.12,y+.36,.48,.05,c)
box(s,9.1,3.45,3.25,2.65,PANEL,line=ORANGE)
pill(s,9.45,3.77,1.5,"Tool action",fill=PANEL2,color=ORANGE)
txt(s,9.45,4.42,2.52,.35,"create_preparation_schedule()",12,TEXT,True)
txt(s,9.45,5.12,2.52,.42,"Outputs an importable calendar—not another paragraph.",11,MUTED)
txt(s,9.35,6.42,2.8,.2,"demo_2_single_agent.py",9,MUTED,True,align=PP_ALIGN.RIGHT)

# 16 — Split threshold
s=prs.slides.add_slide(blank); header(s,"Split the system only after complexity earns it","Complexity threshold",16)
connector(s,1.15,5.65,11.9,5.65,MUTED,1.8); connector(s,1.15,5.65,1.15,1.62,MUTED,1.8)
txt(s,.62,1.72,.32,3.2,"QUALITY",8,MUTED,True)
txt(s,9.92,6.04,2.0,.2,"TASK COMPLEXITY →",9,MUTED,True)
# stair curve
points=[(1.35,5.2),(3.15,4.75),(4.65,4.05),(6.25,3.75),(7.65,2.75),(9.45,2.25),(11.5,1.75)]
for i in range(len(points)-1): connector(s,*points[i],*points[i+1],CYAN,3)
for i,(x,y) in enumerate(points): circle(s,x-.1,y-.1,.2,CYAN)
connector(s,6.75,1.72,6.75,5.65,ORANGE,1.3)
pill(s,5.88,1.35,1.75,"Split here?",fill=PANEL2,color=ORANGE,line=ORANGE)
signals=[("Distinct expertise","different prompts / tools"),("Scoped context","less interference"),("Parallel work","latency gain"),("Independent control","permissions / evals")]
for i,(t,b) in enumerate(signals):
    x=7.35+(i%2)*2.65; y=3.25+(i//2)*1.05
    node(s,x,y,2.35,.78,t,b,accent=[PURPLE,BLUE,ORANGE,GREEN][i])
txt(s,1.48,5.93,5.35,.3,"Single agent zone: simpler, cheaper, easier to debug",11,CYAN,True)

# 17 — collaboration patterns
s=prs.slides.add_slide(blank); header(s,"Three collaboration patterns cover most multi-agent systems","Patterns",17)
# Manager worker
box(s,.6,1.63,3.88,4.85,PANEL,line=LINE); pill(s,.9,1.92,2.0,"Manager → workers",fill=PANEL2,color=CYAN)
node(s,1.55,2.55,2.0,.75,"MANAGER","decompose + integrate",accent=CYAN)
for i,(t,c) in enumerate([("A",BLUE),("B",PURPLE),("C",ORANGE)]):
    x=.92+i*1.13; node(s,x,4.15,.9,.68,t,"specialist",accent=c); connector(s,2.55,3.3,x+.45,4.15,LINE,1.2)
txt(s,.95,5.35,3.15,.55,"Best for independent specialist tasks with one accountable coordinator.",10,MUTED)
# pipeline
box(s,4.72,1.63,3.88,4.85,PANEL,line=LINE); pill(s,5.02,1.92,1.75,"Pipeline",fill=PANEL2,color=PURPLE)
for i,(t,c) in enumerate([("RESEARCH",BLUE),("ANALYSE",PURPLE),("COACH",CYAN),("DECIDE",ORANGE)]):
    y=2.55+i*.73; node(s,5.2,y,2.45,.55,t,"handoff",accent=c)
    if i<3: chevron(s,6.27,y+.57,.27,.25,c,90)
txt(s,5.02,5.82,3.15,.35,"Best when each role enriches the next.",10,MUTED)
# parallel review
box(s,8.84,1.63,3.88,4.85,PANEL,line=LINE); pill(s,9.14,1.92,2.0,"Parallel + review",fill=PANEL2,color=ORANGE)
for i,(t,c) in enumerate([("VIEW 1",BLUE),("VIEW 2",PURPLE),("VIEW 3",PINK)]): node(s,9.15,2.62+i*.82,1.25,.56,t,"",accent=c)
node(s,10.95,3.17,1.35,1.08,"REVIEWER","reconcile",accent=ORANGE)
for i in range(3): connector(s,10.4,2.9+i*.82,10.95,3.72,LINE,1.1)
txt(s,9.15,5.35,3.05,.55,"Best when diversity and verification improve confidence.",10,MUTED)

# 18 — Demo 3 pipeline
s=prs.slides.add_slide(blank); header(s,"Demo 3 — Placement intelligence as a specialist pipeline","Multi-agent",18)
roles=[("01","JOB RESEARCHER","live market signals",BLUE),("02","RESUME ANALYST","evidence + gaps",PURPLE),("03","SKILLS COACH","14-day milestones",CYAN),("04","INTERVIEWER","targeted questions",PINK),("05","CAREER MANAGER","apply / prepare decision",ORANGE)]
for i,(n,t,b,c) in enumerate(roles):
    x=.48+i*2.56
    box(s,x,2.0,2.15,2.72,PANEL,line=c,lw=1.2)
    txt(s,x+.18,2.2,.38,.22,n,9,c,True)
    circle(s,x+.7,2.58,.75,c,t.split()[0][0],15,BG)
    txt(s,x+.15,3.55,1.85,.42,t,10.2,TEXT,True,align=PP_ALIGN.CENTER)
    txt(s,x+.18,4.17,1.78,.32,b,8.8,MUTED,align=PP_ALIGN.CENTER)
    if i<4: chevron(s,x+2.2,3.05,.25,.36,c)
box(s,1.62,5.36,10.1,.85,PANEL2,line=LINE)
rich(s,1.9,5.62,9.55,.3,[("Shared objective: ",CYAN,True),("turn one resume and one job description into an evidence-based placement strategy.",TEXT,False)],13,align=PP_ALIGN.CENTER)
txt(s,9.6,6.46,2.5,.2,"demo_3_multi_agent.py",9,MUTED,True,align=PP_ALIGN.RIGHT)

# 19 — Agent contracts
s=prs.slides.add_slide(blank); header(s,"Multi-agent quality depends on contracts—not agent count","Coordination",19)
box(s,.65,1.62,5.6,4.95,PANEL,line=LINE)
pill(s,.95,1.92,1.62,"Message contract",fill=PANEL2,color=CYAN)
contract=[("TASK","one bounded objective"),("INPUT","minimum relevant context"),("OUTPUT","schema + evidence"),("LIMITS","time • tokens • tools"),("STATUS","done • blocked • escalate")]
for i,(k,v) in enumerate(contract):
    y=2.58+i*.68
    pill(s,.95,y,1.15,k,fill=PANEL2,color=[BLUE,PURPLE,CYAN,ORANGE,RED][i])
    txt(s,2.35,y+.07,3.4,.22,v,11,TEXT)
box(s,6.75,1.62,5.9,4.95,PANEL,line=LINE)
pill(s,7.05,1.92,1.48,"Shared state",fill=PANEL2,color=PURPLE)
for i,(t,b,c) in enumerate([("BLACKBOARD","agents read/write shared facts",BLUE),("HANDOFF","output becomes next input",PURPLE),("EVENT BUS","agents react to events",CYAN),("MANAGER STATE","coordinator owns progress",ORANGE)]):
    y=2.55+i*.82; node(s,7.05,y,5.0,.62,t,b,accent=c)
txt(s,7.05,5.94,5.0,.3,"Prefer explicit state over hidden conversational assumptions.",10,MUTED,True)

# 20 — comparison matrix
s=prs.slides.add_slide(blank); header(s,"Single versus multi-agent: architecture is a trade-off","Decision matrix",20)
headers=[("DIMENSION",.65,2.55,MUTED),("SINGLE AGENT",3.32,4.1,CYAN),("MULTI-AGENT",7.55,5.1,ORANGE)]
for t,x,w,c in headers:
    box(s,x,1.62,w,.6,PANEL2,line=LINE); txt(s,x+.2,1.81,w-.4,.2,t,9.5,c,True)
rows=[("Best fit","coherent task","separable specialist work"),("Context","shared","scoped per role"),("Coordination","minimal","protocol required"),("Latency","serial","parallel or pipeline"),("Cost","usually lower","usually higher"),("Debugging","one trace","distributed traces"),("Permissions","one boundary","role-based boundaries")]
for i,(d,a,b) in enumerate(rows):
    y=2.34+i*.58; fill=PANEL if i%2==0 else BG
    for x,w in [(.65,2.55),(3.32,4.1),(7.55,5.1)]: box(s,x,y,w,.5,fill,line=LINE,radius=False)
    txt(s,.85,y+.14,2.1,.2,d,9.5,MUTED,True); txt(s,3.52,y+.14,3.65,.2,a,10,TEXT); txt(s,7.75,y+.14,4.7,.2,b,10,TEXT)
pill(s,4.5,6.63,4.35,"Start simple. Earn complexity.",fill=PANEL2,color=CYAN,line=CYAN)

# 21 — decision tree
s=prs.slides.add_slide(blank); header(s,"A five-question architecture test","Decision framework",21)
questions=[("Can one agent solve it reliably?",CYAN),("Are subtasks genuinely separable?",BLUE),("Do roles need distinct tools or permissions?",PURPLE),("Does parallelism or review improve outcomes?",ORANGE),("Is the gain worth cost and coordination?",GREEN)]
for i,(q,c) in enumerate(questions):
    x=.75+i*2.48; y=2.15+(i%2)*1.12
    box(s,x,y,2.12,.82,PANEL,line=c)
    circle(s,x+.15,y+.2,.42,c,str(i+1),10,BG)
    txt(s,x+.7,y+.14,1.25,.48,q,9.5,TEXT,True)
    if i<4: chevron(s,x+2.17,y+.25,.23,.3,c)
box(s,1.02,4.8,5.15,1.15,PANEL2,line=CYAN)
txt(s,1.35,5.05,4.5,.26,"NO appears early",11,CYAN,True,align=PP_ALIGN.CENTER)
txt(s,1.35,5.48,4.5,.24,"single agent or deterministic workflow",10,TEXT,align=PP_ALIGN.CENTER)
box(s,7.15,4.8,5.15,1.15,PANEL2,line=ORANGE)
txt(s,7.48,5.05,4.5,.26,"YES survives all five",11,ORANGE,True,align=PP_ALIGN.CENTER)
txt(s,7.48,5.48,4.5,.24,"multi-agent may create measurable value",10,TEXT,align=PP_ALIGN.CENTER)

# 22 — production reference
s=prs.slides.add_slide(blank); header(s,"Production architecture separates experience, intelligence and control","Reference architecture",22)
layers=[("EXPERIENCE","web • mobile • chat • API",BLUE),("ORCHESTRATION","routing • state • retries • approvals",CYAN),("INTELLIGENCE","Gemini • prompts • retrieval • memory",PURPLE),("ACTION","functions • services • databases",ORANGE)]
for i,(t,b,c) in enumerate(layers):
    y=1.55+i*1.08
    box(s,.78,y,8.42,.82,PANEL,line=c)
    box(s,.78,y,2.0,.82,c,line=c)
    txt(s,.95,y+.26,1.66,.2,t,9.5,BG,True,align=PP_ALIGN.CENTER)
    txt(s,3.12,y+.24,5.55,.24,b,12,TEXT,True)
box(s,9.72,1.55,2.78,4.06,PANEL2,line=RED)
pill(s,10.14,1.89,1.95,"Control plane",fill=PANEL,color=RED,line=LINE)
controls=[("IDENTITY","who"),("POLICY","may do"),("OBSERVE","what happened"),("EVALS","was it good"),("COST","was it worth it")]
for i,(t,b) in enumerate(controls):
    y=2.55+i*.54; txt(s,10.02,y,1.15,.2,t,8.5,RED,True); txt(s,11.22,y,1.0,.2,b,9.5,TEXT)
box(s,.78,6.12,11.72,.48,PANEL,line=LINE)
txt(s,1.0,6.25,11.28,.2,"Trace every decision across model calls, tool calls, state transitions and human approvals.",10,MUTED,True,align=PP_ALIGN.CENTER)

# 23 — guardrails
s=prs.slides.add_slide(blank); header(s,"Autonomy becomes trustworthy when every action is bounded","Guardrails",23)
circle(s,5.25,1.72,2.84,PANEL2,"CONTROLLED\nAUTONOMY",17,TEXT,line=CYAN)
guards=[(.65,1.65,"LEAST PRIVILEGE","only required access",CYAN),(1.0,4.72,"HUMAN APPROVAL","before commitments",ORANGE),(9.98,1.65,"BOUNDED RUNS","time • spend • steps",PURPLE),(9.62,4.72,"SAFE FALLBACK","retry • degrade • escalate",RED)]
for x,y,t,b,c in guards:
    node(s,x,y,2.7,1.0,t,b,accent=c)
    connector(s,x+2.7 if x<5 else 8.09,y+.5,5.25 if x<5 else x,y+.5,c,1.4)
pill(s,4.63,5.4,4.05,"Validate inputs + outputs",fill=PANEL,color=GREEN,line=GREEN)
txt(s,3.4,6.27,6.55,.32,"High impact + low reversibility = mandatory human checkpoint",11,MUTED,True,align=PP_ALIGN.CENTER)

# 24 — evaluation dashboard
s=prs.slides.add_slide(blank); header(s,"Evaluate the system—not just the final paragraph","Agent evaluation",24)
metrics=[("92%","TASK SUCCESS",CYAN),("84%","TOOL ACCURACY",BLUE),("1.8","AVG. RETRIES",ORANGE),("4.2s","P95 LATENCY",PURPLE),("₹0.42","COST / RUN",GREEN)]
for i,(v,l,c) in enumerate(metrics): metric(s,.62+i*2.48,1.62,v,l,color=c,w=2.15)
box(s,.62,3.18,5.8,2.72,PANEL,line=LINE)
pill(s,.92,3.48,1.52,"Offline evals",fill=PANEL2,color=BLUE)
txt(s,.92,4.08,4.95,1.25,"• golden tasks\n• tool-selection tests\n• policy simulations\n• regression thresholds",12,TEXT)
box(s,6.82,3.18,5.85,2.72,PANEL,line=LINE)
pill(s,7.12,3.48,1.52,"Online signals",fill=PANEL2,color=CYAN)
txt(s,7.12,4.08,4.95,1.25,"• completion + abandonment\n• human overrides\n• cost + latency\n• downstream outcome",12,TEXT)
txt(s,.75,6.38,11.8,.24,"For placement demos: evidence accuracy • plan feasibility • interview relevance • calendar completion",10,MUTED,True,align=PP_ALIGN.CENTER)

# 25 — runbook
s=prs.slides.add_slide(blank); header(s,"Live demo runbook: one storyline, increasing agency","Demo sequence",25)
demos=[("01","LLM CALL","Gap analysis","python demo_1_llm_call.py",BLUE,"2 min"),(
"02","SINGLE AGENT","Plan + calendar","python demo_2_single_agent.py --days 14",CYAN,"4 min"),(
"03","MULTI-AGENT","Placement intelligence","python demo_3_multi_agent.py",ORANGE,"6 min")]
for i,(n,t,b,cmd,c,time_) in enumerate(demos):
    y=1.55+i*1.58
    box(s,.65,y,12.02,1.28,PANEL,line=c,lw=1.2)
    circle(s,.92,y+.34,.56,c,n,11,BG)
    txt(s,1.75,y+.22,2.1,.22,t,11.5,c,True)
    txt(s,1.75,y+.65,2.2,.25,b,12,TEXT,True)
    box(s,4.25,y+.25,6.65,.72,BG,line=LINE)
    txt(s,4.52,y+.47,6.1,.2,cmd,10.5,TEXT,False,font="Menlo")
    pill(s,11.35,y+.43,.9,time_,fill=PANEL2,color=c)
box(s,2.05,6.35,9.2,.44,PANEL2,line=LINE)
txt(s,2.25,6.47,8.8,.18,"Narrate the architecture while the output streams—do not narrate the prompt.",9.5,MUTED,True,align=PP_ALIGN.CENTER)

# 26 — session map
s=prs.slides.add_slide(blank); header(s,"A 60-minute session with three deliberate peaks","Facilitation",26)
segments=[("0–5","HOOK","everyday AI",BLUE,1.0),(
"5–15","MODEL","AI → agents",PURPLE,2.35),("15–27","SINGLE","loop + demo",CYAN,4.15),("27–43","MULTI","patterns + demo",ORANGE,6.5),("43–53","PRODUCTION","guardrails + evals",GREEN,9.05),("53–60","DESIGN","audience challenge",PINK,11.35)]
connector(s,.92,3.65,12.42,3.65,LINE,3)
for time_,t,b,c,x in segments:
    circle(s,x,3.38,.54,c)
    txt(s,x-.38,2.42,1.3,.24,time_,10,c,True,align=PP_ALIGN.CENTER)
    txt(s,x-.62,4.18,1.8,.24,t,10.5,TEXT,True,align=PP_ALIGN.CENTER)
    txt(s,x-.67,4.62,1.9,.4,b,9,MUTED,align=PP_ALIGN.CENTER)
box(s,1.15,5.75,11.05,.62,PANEL2,line=LINE)
txt(s,1.4,5.94,10.55,.25,"Audience prompt: “What daily problem around you deserves an agent—and what must remain human?”",12,CYAN,True,align=PP_ALIGN.CENTER)

# 27 — closing
s=prs.slides.add_slide(blank); canvas(s,number=27)
pill(s,.65,.62,1.48,"Takeaway",fill=PANEL2,color=CYAN)
txt(s,.65,1.55,12.0,.55,"AI can generate.",29,MUTED,True,align=PP_ALIGN.CENTER)
txt(s,.65,2.48,12.0,.62,"Agents can act.",37,CYAN,True,align=PP_ALIGN.CENTER)
txt(s,.65,3.58,12.0,.7,"Multi-agent systems can collaborate.",36,ORANGE,True,align=PP_ALIGN.CENTER)
connector(s,3.05,4.78,10.28,4.78,LINE,1.2)
txt(s,2.05,5.27,9.25,.68,"Start with the outcome. Use the simplest architecture that solves it reliably—and keep control explicit.",17,TEXT,True,align=PP_ALIGN.CENTER)
pill(s,4.62,6.35,4.1,"What will you turn into an agent?",fill=PANEL,color=PINK,line=PINK)

# 28 — sources
s=prs.slides.add_slide(blank); header(s,"References used for implementation details","Sources",28)
refs=[("Gemini API — Getting started","ai.google.dev/gemini-api/docs/get-started",CYAN),(
"Gemini API — Function calling","ai.google.dev/gemini-api/docs/function-calling",PURPLE),(
"Gemini model documentation","ai.google.dev/gemini-api/docs/models",BLUE),(
"Gemini API key guidance","ai.google.dev/gemini-api/docs/api-key",ORANGE),(
"Local placement demos","/Users/ritishjogi/Downloads/st-peters",GREEN)]
for i,(t,u,c) in enumerate(refs):
    y=1.62+i*.92
    box(s,1.05,y,11.15,.72,PANEL,line=LINE)
    circle(s,1.3,y+.18,.36,c,str(i+1),8,BG)
    txt(s,1.9,y+.14,3.45,.22,t,11,TEXT,True)
    txt(s,5.3,y+.14,6.45,.25,u,10,c)
txt(s,1.05,6.45,11.15,.25,"Product capabilities evolve; verify model availability and API behavior before the live session.",9,MUTED,True,align=PP_ALIGN.CENTER)


prs.core_properties.title = "From AI to Agents: Building Intelligent Systems from Single to Multi-Agent"
prs.core_properties.subject = "Advanced dark-theme session deck with Gemini placement demos"
prs.core_properties.author = "Ritish Jogi"
prs.core_properties.keywords = "AI agents, Gemini API, single agent, multi-agent, placement intelligence"
prs.save(OUT)
print(OUT)
print(f"slides={len(prs.slides)}")
